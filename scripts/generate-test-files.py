from __future__ import annotations

import csv
import json
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test-files"
BURGUNDY = "4F222C"
BURGUNDY_LIGHT = "F5ECEE"
CREAM = "F8F5EF"
GOLD = "C49A55"
INK = "262129"
MUTED = "6F6870"
GREEN = "2D6A4F"
RED = "A63D40"


TRANSACTIONS = [
    ("TX-1001", "2026-08-02", "Northstar Office Supply", 428.16, "Operations", "Printer paper and toner", "NS-8841", "Corporate card", "Jordan Lee", "6100", "Chicago", "PO-260731", "Matched", "Low", "Q3-OPS"),
    ("TX-1002", "2026-08-03", "Cloud Harbor", 1240.00, "Engineering", "August cloud hosting", "CH-2026-08", "ACH", "Priya Shah", "6420", "Remote", "PO-260612", "Matched", "Low", "PLATFORM"),
    ("TX-1003", "2026-08-04", "Metro Rail", 86.40, "Sales", "Client meeting travel", "MR-3019", "Corporate card", "Elena Ruiz", "7310", "New York", "", "Receipt only", "Low", "NORTHEAST"),
    ("TX-1004", "2026-08-05", "Brightline Catering", 612.75, "People", "Quarterly team lunch", "BC-7712", "Corporate card", "Tessa Cole", "6210", "Boston", "PO-260804", "Matched", "Low", "PEOPLE-Q3"),
    ("TX-1005", "2026-08-06", "Acme Industrial Parts", 4800.00, "Facilities", "Replacement pump assembly", "AI-4408", "ACH", "Marcus Bell", "6810", "Chicago", "PO-260755", "Matched", "Medium", "FAC-CHILLER"),
    ("TX-1006", "2026-08-08", "Acme Industrial Parts", 4800.00, "Facilities", "Replacement pump assembly", "AI-4418", "ACH", "Marcus Bell", "6810", "Chicago", "PO-260755", "Exception", "High", "FAC-CHILLER"),
    ("TX-1007", "2026-08-08", "Redwood Advisory", 4950.00, "Strategy", "Market assessment phase one", "RA-901", "Wire", "Dana Wu", "6650", "Corporate", "", "No PO", "Medium", "GROWTH-26"),
    ("TX-1008", "2026-08-08", "Redwood Advisory", 4900.00, "Strategy", "Market assessment phase two", "RA-902", "Wire", "Dana Wu", "6650", "Corporate", "", "No PO", "Medium", "GROWTH-26"),
    ("TX-1009", "2026-08-08", "Redwood Advisory", 4850.00, "Strategy", "Market assessment phase three", "RA-903", "Wire", "Dana Wu", "6650", "Corporate", "", "No PO", "Medium", "GROWTH-26"),
    ("TX-1010", "2026-08-09", "Night Owl Liquor", 3275.00, "Engineering", "Cloud infrastructure migration", "NO-114", "Corporate card", "Morgan Hale", "6420", "Remote", "", "Receipt missing", "High", "MIGRATION"),
    ("TX-1011", "2026-08-10", "Apex Strategic Services", 18750.00, "Executive", "Professional services rendered", "AP-2026-18", "Wire", "Cameron Price", "6650", "Corporate", "", "No PO", "High", "CEO-OFFICE"),
    ("TX-1012", "2026-08-11", "Greenway Couriers", 144.30, "Legal", "Document delivery", "GC-5583", "Corporate card", "Noah Kim", "6290", "New York", "", "Receipt only", "Low", "LEGAL-OPS"),
    ("TX-1013", "2026-08-12", "Prism Software", 2799.00, "Design", "Annual prototyping software", "PS-9931", "ACH", "Maya Patel", "6430", "Remote", "PO-260420", "Matched", "Low", "DESIGN-SYS"),
    ("TX-1014", "2026-08-13", "Harbor Hotel", 892.46, "Sales", "Customer conference lodging", "HH-8204", "Corporate card", "Elena Ruiz", "7310", "New York", "", "Receipt only", "Low", "NORTHEAST"),
    ("TX-1015", "2026-08-14", "Northstar Office Supply", 91.12, "Operations", "Shipping labels", "NS-8902", "Corporate card", "Jordan Lee", "6100", "Chicago", "", "Receipt only", "Low", "Q3-OPS"),
    ("TX-1016", "2026-08-15", "Quick Cash Services", 9200.00, "Marketing", "Miscellaneous campaign support", "QCS-77", "Wire", "Avery Brooks", "7240", "Boston", "", "No PO", "High", "FALL-LAUNCH"),
    ("TX-1017", "2026-08-17", "Lumen Research", 6400.00, "Product", "Customer interviews and analysis", "LR-216", "ACH", "Isaac Chen", "6640", "Remote", "PO-260701", "Matched", "Low", "DISCOVERY"),
    ("TX-1018", "2026-08-18", "City Electric", 1328.44, "Facilities", "Server room electrical work", "CE-8041", "Check", "Marcus Bell", "6810", "Chicago", "PO-260779", "Matched", "Low", "FAC-SERVER"),
    ("TX-1019", "2026-08-19", "Cloud Harbor", 1240.00, "Engineering", "September cloud hosting", "CH-2026-09", "ACH", "Priya Shah", "6420", "Remote", "PO-260612", "Matched", "Low", "PLATFORM"),
    ("TX-1020", "2026-08-20", "Evergreen Printing", 538.20, "Marketing", "Event brochures", "EP-6610", "Corporate card", "Avery Brooks", "7240", "Boston", "PO-260818", "Matched", "Low", "FALL-LAUNCH"),
    ("TX-1021", "2026-08-21", "Apex Strategic Services", 18750.00, "Executive", "Professional services rendered", "AP-2026-19", "Wire", "Cameron Price", "6650", "Corporate", "", "No PO", "High", "CEO-OFFICE"),
    ("TX-1022", "2026-08-22", "Sunset Grocery", 2480.00, "Information Security", "Penetration testing services", "SG-940", "Corporate card", "Morgan Hale", "6460", "Remote", "", "Receipt missing", "High", "SEC-AUDIT"),
    ("TX-1023", "2026-08-23", "Atlas Telecom", 780.32, "Operations", "Office internet service", "AT-202608", "ACH", "Jordan Lee", "6410", "Chicago", "PO-260101", "Matched", "Low", "Q3-OPS"),
    ("TX-1024", "2026-08-24", "Clearwater Insurance", 3520.00, "Finance", "Quarterly liability coverage", "CI-3118", "ACH", "Rina Fox", "6520", "Corporate", "PO-260211", "Matched", "Low", "RISK-26"),
    ("TX-1025", "2026-08-25", "Metro Rail", 64.20, "Sales", "Regional customer visit", "MR-3092", "Corporate card", "Elena Ruiz", "7310", "New York", "", "Receipt only", "Low", "NORTHEAST"),
    ("TX-1026", "2026-08-26", "Bluebird Air", 684.18, "Sales", "Customer summit airfare", "BA-94408", "Corporate card", "Elena Ruiz", "7310", "New York", "", "Receipt only", "Low", "NORTHEAST"),
    ("TX-1027", "2026-08-26", "Juniper Labs", 3300.00, "Product", "Prototype usability study", "JL-1847", "ACH", "Isaac Chen", "6640", "Remote", "PO-260825", "Matched", "Low", "DISCOVERY"),
    ("TX-1028", "2026-08-27", "Civic Parking", 38.00, "Legal", "Court filing visit parking", "CP-8182", "Corporate card", "Noah Kim", "7310", "New York", "", "Receipt only", "Low", "LEGAL-OPS"),
    ("TX-1029", "2026-08-27", "Kite & Co. Design", 2140.00, "Marketing", "Campaign illustration package", "KC-4811", "ACH", "Avery Brooks", "7240", "Boston", "PO-260822", "Matched", "Low", "FALL-LAUNCH"),
    ("TX-1030", "2026-08-28", "Northstar Office Supply", 376.89, "Operations", "Facilities restock", "NS-9021", "Corporate card", "Jordan Lee", "6100", "Chicago", "PO-260901", "Matched", "Low", "Q3-OPS"),
    ("TX-1031", "2026-08-29", "Harbor Hotel", 1044.60, "Strategy", "Annual planning offsite lodging", "HH-8377", "Corporate card", "Dana Wu", "7340", "Corporate", "", "Receipt only", "Low", "PLAN-27"),
    ("TX-1032", "2026-08-30", "Brightline Catering", 458.10, "Engineering", "Incident review working dinner", "BC-7891", "Corporate card", "Priya Shah", "6210", "Remote", "", "Receipt only", "Low", "PLATFORM"),
    ("TX-1033", "2026-09-01", "Atlas Telecom", 780.32, "Operations", "Office internet service", "AT-202609", "ACH", "Jordan Lee", "6410", "Chicago", "PO-260101", "Matched", "Low", "Q3-OPS"),
    ("TX-1034", "2026-09-02", "Cedar Backgrounds", 1185.00, "People", "Pre-employment screening", "CB-2290", "ACH", "Tessa Cole", "6230", "Boston", "PO-260220", "Matched", "Low", "PEOPLE-Q3"),
    ("TX-1035", "2026-09-03", "Cloud Harbor", 286.41, "Engineering", "Usage overage", "CH-2026-09A", "ACH", "Priya Shah", "6420", "Remote", "PO-260612", "Matched", "Low", "PLATFORM"),
    ("TX-1036", "2026-09-04", "Westlake Hardware", 267.83, "Facilities", "Door hardware and fasteners", "WH-77218", "Corporate card", "Marcus Bell", "6810", "Chicago", "", "Receipt only", "Low", "FAC-GENERAL"),
    ("TX-1037", "2026-09-05", "Lumen Research", 3200.00, "Product", "Interview transcription milestone", "LR-223", "ACH", "Isaac Chen", "6640", "Remote", "PO-260701", "Matched", "Low", "DISCOVERY"),
    ("TX-1038", "2026-09-06", "Evergreen Printing", 189.75, "Marketing", "Name badges and signage", "EP-6692", "Corporate card", "Avery Brooks", "7240", "Boston", "", "Receipt only", "Low", "FALL-LAUNCH"),
    ("TX-1039", "2026-09-07", "Greenway Couriers", 176.55, "Legal", "Regulatory filing delivery", "GC-5714", "Corporate card", "Noah Kim", "6290", "New York", "", "Receipt only", "Low", "LEGAL-OPS"),
    ("TX-1040", "2026-09-08", "Prism Software", 420.00, "Design", "Additional editor seats", "PS-10018", "Corporate card", "Maya Patel", "6430", "Remote", "PO-260420", "Matched", "Low", "DESIGN-SYS"),
]

HEADERS = ["transaction_id", "posting_date", "vendor", "gross_amount", "department", "description", "invoice_id", "payment_method", "requestor", "gl_code", "office", "purchase_order", "match_status", "review_tier", "project_code"]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_docx() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)

    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(4.9)
    header.columns[1].width = Inches(2.25)
    left, right = header.rows[0].cells
    set_cell_shading(left, BURGUNDY)
    set_cell_shading(right, BURGUNDY)
    set_cell_margins(left, 180, 220, 180, 220)
    set_cell_margins(right, 180, 180, 180, 220)
    p = left.paragraphs[0]
    r = p.add_run("NORTHLIGHT / FINANCE")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(255, 255, 255)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("EXPENSE REPORT\nER-8821")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(255, 255, 255)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("August 2026 expense reconciliation")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(BURGUNDY)
    sub = doc.add_paragraph("Submitted for manager and Finance review · Generated September 2, 2026")
    sub.paragraph_format.space_after = Pt(12)
    sub.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    info = doc.add_table(rows=2, cols=4)
    info.autofit = False
    info_data = [
        ("EMPLOYEE", "Morgan Hale", "DEPARTMENT", "Information Security"),
        ("EMPLOYEE ID", "NL-0472", "COST CENTER", "6460 · Security Operations"),
    ]
    for row_i, values in enumerate(info_data):
        for col_i, value in enumerate(values):
            c = info.cell(row_i, col_i)
            set_cell_shading(c, CREAM if row_i == 0 else "FFFFFF")
            set_cell_margins(c, 90, 110, 90, 110)
            p = c.paragraphs[0]
            r = p.add_run(value)
            if col_i % 2 == 0:
                r.bold = True
                r.font.size = Pt(7)
                r.font.color.rgb = RGBColor.from_string(MUTED)
            else:
                r.bold = True
                r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    summary = doc.add_table(rows=1, cols=4)
    for i, (label, value) in enumerate((("TOTAL CLAIMED", "$20,032.57"), ("CORPORATE CARD", "$19,666.64"), ("OUT OF POCKET", "$365.93"), ("MISSING RECEIPTS", "3"))):
        c = summary.cell(0, i)
        set_cell_shading(c, BURGUNDY_LIGHT)
        set_cell_margins(c, 100, 110, 100, 110)
        p = c.paragraphs[0]
        p.add_run(label + "\n").font.size = Pt(6.5)
        v = p.add_run(value)
        v.bold = True
        v.font.size = Pt(12)
        v.font.color.rgb = RGBColor.from_string(BURGUNDY)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(6)
    r = heading.add_run("EXPENSE DETAIL")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(BURGUNDY)

    rows = [
        ("Aug 03", "Cloud Harbor", "CH-2026-08", "Cloud hosting subscription", "Card · 1842", "$1,240.00", "Attached"),
        ("Aug 06", "Cipher Conference", "CC-2608", "Security leadership registration", "Card · 1842", "$1,895.00", "Attached"),
        ("Aug 09", "Night Owl Liquor", "NO-114", "Cloud infrastructure migration", "Card · 1842", "$3,275.00", "Missing"),
        ("Aug 11", "Bluebird Air", "BA-94011", "Travel to security summit", "Card · 1842", "$684.18", "Attached"),
        ("Aug 12", "Harbor Hotel", "HH-8208", "Security summit lodging", "Card · 1842", "$892.46", "Attached"),
        ("Aug 13", "Metro Rail", "MR-3067", "Airport transfer", "Personal", "$64.20", "Attached"),
        ("Aug 14", "Juniper Cafe", "JC-1882", "Client security review lunch", "Personal", "$126.75", "Attached"),
        ("Aug 16", "Sunset Grocery", "SG-940", "Penetration testing services", "Card · 1842", "$2,480.00", "Missing"),
        ("Aug 18", "Quick Cash Services", "QCS-77", "Miscellaneous support", "Card · 1842", "$9,200.00", "Missing"),
        ("Aug 20", "Metro Rail", "MR-3092", "Regional office visit", "Personal", "$64.20", "Attached"),
        ("Aug 22", "Civic Parking", "CP-8810", "Vendor assessment meeting", "Personal", "$38.00", "Attached"),
        ("Aug 24", "Greenway Couriers", "GC-5624", "Signed audit materials", "Personal", "$72.78", "Attached"),
    ]
    widths = [0.53, 1.15, 0.73, 1.75, 0.82, 0.72, 0.63]
    table = doc.add_table(rows=1, cols=7)
    table.autofit = False
    for i, (label, width) in enumerate(zip(("DATE", "MERCHANT", "REFERENCE", "BUSINESS PURPOSE", "PAID VIA", "AMOUNT", "RECEIPT"), widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        set_cell_shading(cell, BURGUNDY)
        set_cell_margins(cell, 90, 70, 90, 70)
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(6.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            c = cells[i]
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(c, "FFFFFF" if row_i % 2 == 0 else CREAM)
            set_cell_margins(c, 75, 70, 75, 70)
            p = c.paragraphs[0]
            if i == 5:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(value)
            r.font.size = Pt(7.2)
            if value == "Missing":
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(RED)

    note = doc.add_table(rows=1, cols=2)
    note.autofit = False
    note.columns[0].width = Inches(5.05)
    note.columns[1].width = Inches(2.1)
    c = note.cell(0, 0)
    set_cell_shading(c, "FFF8E8")
    set_cell_margins(c, 110, 120, 110, 120)
    p = c.paragraphs[0]
    r = p.add_run("SUBMITTER NOTE\n")
    r.bold = True
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    p.add_run("The three expenses marked missing were approved verbally during the migration incident. Replacement receipts have been requested.").font.size = Pt(8)
    c = note.cell(0, 1)
    set_cell_shading(c, CREAM)
    set_cell_margins(c, 110, 120, 110, 120)
    p = c.paragraphs[0]
    p.add_run("MANAGER REVIEW\n").bold = True
    p.add_run("Cameron Price\nApproved 09/02/26\nException noted").font.size = Pt(8)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("NORTHLIGHT INTERNAL · Expense policy FIN-204 · Retain 7 years · ER-8821")
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.save(OUT / "expense-report.docx")


def add_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = BURGUNDY
    ws.oddHeader.center.text = "&BNORTHLIGHT / FINANCE OPERATIONS&B"
    ws.oddHeader.right.text = "Page &P of &N"
    ws.oddFooter.left.text = "Internal · AP transaction register · Exported Sep 9, 2026"
    ws.oddFooter.right.text = "US01 · USD"
    for col, header in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=header)
        c.font = Font(size=8, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BURGUNDY)
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    for row_i, row in enumerate(TRANSACTIONS, 2):
        for col_i, value in enumerate(row, 1):
            c = ws.cell(row=row_i, column=col_i, value=value)
            c.font = Font(name="Aptos", size=9, color=INK)
            c.fill = PatternFill("solid", fgColor="FFFFFF" if row_i % 2 else CREAM)
            c.alignment = Alignment(vertical="center")
            c.border = Border(bottom=Side(style="hair", color="E6DFD7"))
        ws.cell(row=row_i, column=4).number_format = '$#,##0.00;[Red]-$#,##0.00'
        risk_cell = ws.cell(row=row_i, column=14)
        if risk_cell.value == "High":
            risk_cell.fill = PatternFill("solid", fgColor="F7D9DA")
            risk_cell.font = Font(size=9, bold=True, color=RED)
        elif risk_cell.value == "Medium":
            risk_cell.fill = PatternFill("solid", fgColor="FBECCB")
    widths = [13, 13, 26, 14, 20, 34, 17, 18, 18, 11, 14, 16, 15, 12, 16]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    table_ref = f"A1:O{1 + len(TRANSACTIONS)}"
    transactions_table = XLTable(displayName="NorthlightTransactions", ref=table_ref)
    transactions_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(transactions_table)
    ws.conditional_formatting.add(f"D2:D{1 + len(TRANSACTIONS)}", ColorScaleRule(start_type="min", start_color="F7F3ED", end_type="max", end_color="C98E98"))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:1"

    vendors = wb.create_sheet("Vendor Directory")
    vendors.sheet_view.showGridLines = False
    vendor_headers = ["Vendor ID", "Legal name", "Category", "Status", "Payment account", "Terms", "Tax region", "Onboarded", "Last review", "Risk owner"]
    vendor_rows = [
        ("V-0184", "Northstar Office Supply LLC", "Office supplies", "Approved", "Business checking · 1842", "Net 30", "US-IL", "2021-04-16", "2026-03-14", "Procurement"),
        ("V-0271", "Cloud Harbor Systems Inc.", "Cloud infrastructure", "Approved", "Business checking · 6210", "Net 30", "US-VA", "2022-10-03", "2026-07-22", "IT"),
        ("V-0318", "Acme Industrial Parts Co.", "Industrial equipment", "Approved", "Business checking · 7784", "Net 45", "US-OH", "2020-02-18", "2026-01-11", "Facilities"),
        ("V-0406", "Redwood Advisory Group LLC", "Management consulting", "Approved", "Business checking · 9927", "Due on receipt", "US-CA", "2025-12-02", "2026-06-30", "Strategy"),
        ("V-0440", "Apex Strategic Services LLC", "General consulting", "Pending review", "Personal checking · 4401", "Due on receipt", "US-DE", "2026-07-29", "Not completed", "Executive"),
        ("V-0447", "Night Owl Liquor Corp.", "Retail liquor store", "Unapproved", "Corporate card", "Immediate", "US-NY", "Not onboarded", "Not reviewed", "Unassigned"),
        ("V-0448", "Sunset Grocery Market", "Grocery and household retail", "Unapproved", "Corporate card", "Immediate", "US-MA", "Not onboarded", "Not reviewed", "Unassigned"),
        ("V-0452", "Quick Cash Services LLC", "Alternative financial services", "Restricted", "Business checking · 1170", "Due on receipt", "US-NJ", "2026-08-12", "Escalated", "Compliance"),
        ("V-0229", "Lumen Research Partners", "Market research", "Approved", "Business checking · 3048", "Net 30", "US-IL", "2023-01-20", "2026-05-18", "Product"),
        ("V-0122", "Atlas Telecom Inc.", "Telecommunications", "Approved", "Business checking · 8305", "Net 30", "US-TX", "2019-06-01", "2026-02-09", "IT"),
    ]
    vendors.merge_cells("A1:J1")
    vendors["A1"] = "VENDOR DIRECTORY  /  CONTROL SNAPSHOT"
    vendors["A1"].font = Font(size=17, bold=True, color="FFFFFF")
    vendors["A1"].fill = PatternFill("solid", fgColor=BURGUNDY)
    vendors.row_dimensions[1].height = 32
    vendors.merge_cells("A2:J2")
    vendors["A2"] = "Includes active and exception vendors referenced by the transaction register"
    vendors["A2"].font = Font(size=10, color=MUTED)
    for col, header in enumerate(vendor_headers, 1):
        c = vendors.cell(row=4, column=col, value=header.upper())
        c.font = Font(size=8, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BURGUNDY)
    for row_i, row in enumerate(vendor_rows, 5):
        for col_i, value in enumerate(row, 1):
            c = vendors.cell(row=row_i, column=col_i, value=value)
            c.font = Font(size=9, color=INK)
            c.fill = PatternFill("solid", fgColor="FFFFFF" if row_i % 2 else CREAM)
        if row[3] in ("Unapproved", "Restricted"):
            vendors.cell(row=row_i, column=4).fill = PatternFill("solid", fgColor="F7D9DA")
            vendors.cell(row=row_i, column=4).font = Font(size=9, bold=True, color=RED)
    for i, width in enumerate([13, 31, 26, 18, 27, 18, 14, 16, 16, 16], 1):
        vendors.column_dimensions[get_column_letter(i)].width = width
    vendors.freeze_panes = "A5"
    vendors.auto_filter.ref = f"A4:J{4 + len(vendor_rows)}"

    readme = wb.create_sheet("Export Notes")
    readme.sheet_view.showGridLines = False
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 76
    readme.merge_cells("A1:B1")
    readme["A1"] = "EXPORT NOTES"
    readme["A1"].font = Font(size=17, bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor=BURGUNDY)
    notes = [("Source system", "Northlight Ledger Cloud / AP subledger"), ("Entity", "Northlight Inc. (US01)"), ("Currency", "USD; gross amount includes sales tax where applicable"), ("Data owner", "Finance Operations / finops@northlight.example"), ("Control note", "Rows marked Exception, No PO, or Receipt missing require secondary review."), ("Confidentiality", "Internal use only. Synthetic training fixture; no real individuals or accounts.")]
    for row_i, (label, value) in enumerate(notes, 3):
        readme.cell(row_i, 1, label).font = Font(size=9, bold=True, color=BURGUNDY)
        readme.cell(row_i, 2, value).font = Font(size=9, color=INK)
        if row_i % 2:
            readme.cell(row_i, 1).fill = PatternFill("solid", fgColor=CREAM)
            readme.cell(row_i, 2).fill = PatternFill("solid", fgColor=CREAM)

    wb.save(OUT / "transactions.xlsx")


def add_csv() -> None:
    with (OUT / "transactions.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        writer.writerows(TRANSACTIONS)


def add_json() -> None:
    records = []
    source = [
        ("V-0184", "Northstar Office Supply LLC", "Northstar Office Supply", "Office supplies", "approved", "Business checking ending 1842", "net_30", "2021-04-16", "2026-03-14", "low", True, "W-9 on file"),
        ("V-0271", "Cloud Harbor Systems Inc.", "Cloud Harbor", "Cloud infrastructure", "approved", "Business checking ending 6210", "net_30", "2022-10-03", "2026-07-22", "low", True, "SOC 2 reviewed"),
        ("V-0318", "Acme Industrial Parts Co.", "Acme Industrial Parts", "Industrial equipment", "approved", "Business checking ending 7784", "net_45", "2020-02-18", "2026-01-11", "medium", True, "Bank update verified by callback"),
        ("V-0406", "Redwood Advisory Group LLC", "Redwood Advisory", "Management consulting", "approved", "Business checking ending 9927", "due_on_receipt", "2025-12-02", "2026-06-30", "medium", False, "SOW renewal pending"),
        ("V-0440", "Apex Strategic Services LLC", "Apex Strategic Services", "General consulting", "pending_review", "Personal checking ending 4401", "due_on_receipt", "2026-07-29", None, "high", False, "Beneficial ownership review incomplete"),
        ("V-0447", "Night Owl Liquor Corp.", "Night Owl Liquor", "Retail liquor store", "unapproved", "Corporate card", "immediate", None, None, "high", False, "Merchant not in vendor portal"),
        ("V-0448", "Sunset Grocery Market", "Sunset Grocery", "Grocery and household retail", "unapproved", "Corporate card", "immediate", None, None, "high", False, "Merchant not in vendor portal"),
        ("V-0452", "Quick Cash Services LLC", "Quick Cash Services", "Alternative financial services", "restricted", "Business checking ending 1170", "due_on_receipt", "2026-08-12", "escalated", "high", False, "Enhanced due diligence required"),
        ("V-0229", "Lumen Research Partners", "Lumen Research", "Market research", "approved", "Business checking ending 3048", "net_30", "2023-01-20", "2026-05-18", "low", True, "Standard annual review"),
        ("V-0122", "Atlas Telecom Inc.", "Atlas Telecom", "Telecommunications", "approved", "Business checking ending 8305", "net_30", "2019-06-01", "2026-02-09", "low", True, "Standard annual review"),
        ("V-0351", "Prism Software Inc.", "Prism Software", "Software subscriptions", "approved", "Business checking ending 5018", "net_30", "2024-05-09", "2026-04-28", "low", True, "Security review current"),
        ("V-0378", "Clearwater Insurance Group", "Clearwater Insurance", "Commercial insurance", "approved", "Business checking ending 2881", "net_30", "2024-11-21", "2026-01-31", "low", True, "Certificate on file"),
    ]
    for vendor_id, legal, display, category, status, account, terms, onboarded, review, risk, po, notes in source:
        records.append({"vendor_id": vendor_id, "legal_name": legal, "display_name": display, "category": category, "status": status, "payment_account": account, "payment_terms": terms, "tax_region": "US", "onboarded_at": onboarded, "last_compliance_review": review, "risk_tier": risk, "purchase_order_required": po, "currency": "USD", "entity_code": "US01", "data_owner": "Finance Operations", "review_notes": notes, "source_system": "Northlight Vendor Portal", "exported_at": "2026-09-09T08:42:16-04:00", "schema_version": "3.2", "synthetic_fixture": True})
    (OUT / "vendor-master.json").write_text(json.dumps(records, indent=2) + "\n", encoding="utf-8")


def add_pdf() -> None:
    path = OUT / "invoice-apex.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.58*inch, leftMargin=0.58*inch, topMargin=0.48*inch, bottomMargin=0.52*inch, title="Apex Strategic Services Invoice AP-2026-18", author="Apex Strategic Services LLC")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body2", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=12, textColor=colors.HexColor(f"#{INK}"))
    small = ParagraphStyle("small", parent=body, fontSize=7, leading=9, textColor=colors.HexColor(f"#{MUTED}"))
    right = ParagraphStyle("right", parent=body, alignment=TA_RIGHT)
    story = []
    brand = Table([[Paragraph("<b>APEX</b><br/><font size='7'>STRATEGIC SERVICES</font>", ParagraphStyle("brand", parent=body, fontSize=20, leading=16, textColor=colors.white)), Paragraph("<b>INVOICE</b><br/><font size='10'>AP-2026-18</font>", ParagraphStyle("invoice", parent=right, fontSize=15, leading=17, textColor=colors.white))]], colWidths=[4.7*inch, 2.15*inch])
    brand.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{BURGUNDY}")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 16), ("RIGHTPADDING", (0,0), (-1,-1), 16), ("TOPPADDING", (0,0), (-1,-1), 13), ("BOTTOMPADDING", (0,0), (-1,-1), 13)]))
    story += [brand, Spacer(1, 0.22*inch)]
    meta = Table([[Paragraph("<font color='#6F6870' size='7'><b>BILL TO</b></font><br/><b>Northlight Inc.</b><br/>155 Wacker Drive, Suite 2400<br/>Chicago, IL 60606<br/>Attn: Office of the CEO", body), Paragraph("<font color='#6F6870' size='7'><b>REMIT TO</b></font><br/><b>Apex Strategic Services LLC</b><br/>1209 Orange Street<br/>Wilmington, DE 19801<br/>ap@apexstrategic.example", body), Paragraph("<font color='#6F6870' size='7'><b>INVOICE DETAILS</b></font><br/>Issue date: <b>Aug 10, 2026</b><br/>Due date: <b>Aug 10, 2026</b><br/>Terms: <b>Due on receipt</b><br/>Currency: <b>USD</b>", body)]], colWidths=[2.42*inch, 2.42*inch, 2.01*inch])
    meta.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{CREAM}")), ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#DED6CD")), ("INNERGRID", (0,0), (-1,-1), 0.5, colors.HexColor("#E7E0D8")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 11), ("RIGHTPADDING", (0,0), (-1,-1), 11), ("TOPPADDING", (0,0), (-1,-1), 10), ("BOTTOMPADDING", (0,0), (-1,-1), 10)]))
    story += [meta, Spacer(1, 0.23*inch), Paragraph("<b>ENGAGEMENT SUMMARY</b>", ParagraphStyle("section", parent=body, fontSize=8, textColor=colors.HexColor(f"#{BURGUNDY}"), spaceAfter=5)), Paragraph("Executive advisory support related to FY27 operating priorities, organizational alignment, and leadership communications. Services provided under engagement letter dated July 29, 2026. Client reference: CEO-OFFICE / Cost center 6650.", body), Spacer(1, 0.18*inch)]
    line_data = [["SERVICE DATE", "DESCRIPTION", "CONSULTANT", "QTY", "RATE", "AMOUNT"], ["Aug 3", "Executive discovery session and stakeholder interviews", "C. Reed", "8.0 hr", "$450.00", "$3,600.00"], ["Aug 4–5", "Operating model assessment and working materials", "C. Reed", "18.0 hr", "$450.00", "$8,100.00"], ["Aug 6", "Leadership alignment workshop preparation", "J. Vale", "10.0 hr", "$375.00", "$3,750.00"], ["Aug 7", "Leadership alignment workshop facilitation", "J. Vale", "6.0 hr", "$375.00", "$2,250.00"], ["Aug 7", "Executive summary and recommended next steps", "C. Reed", "2.0 hr", "$450.00", "$900.00"], ["Aug 7", "Administrative and document production fee", "Project office", "1", "$150.00", "$150.00"]]
    lines = Table(line_data, colWidths=[0.72*inch, 2.74*inch, 0.92*inch, 0.57*inch, 0.76*inch, 0.86*inch], repeatRows=1)
    lines.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor(f"#{BURGUNDY}")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,0), 6.5), ("ALIGN", (3,1), (-1,-1), "RIGHT"), ("FONTNAME", (0,1), (-1,-1), "Helvetica"), ("FONTSIZE", (0,1), (-1,-1), 7.5), ("TEXTCOLOR", (0,1), (-1,-1), colors.HexColor(f"#{INK}")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor(f"#{CREAM}")]), ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#DDD5CC")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6), ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7)]))
    story += [lines, Spacer(1, 0.18*inch)]
    totals = Table([[Paragraph("<font color='#6F6870'>PAYMENT REFERENCE</font><br/><b>AP-2026-18 / Northlight Inc.</b><br/><font size='7'>Please include the invoice number with remittance.</font>", body), "Subtotal", "$18,750.00"], ["", "Tax", "$0.00"], ["", Paragraph("<b>AMOUNT DUE</b>", body), Paragraph("<b>$18,750.00</b>", right)]], colWidths=[4.4*inch, 1.2*inch, 1.25*inch])
    totals.setStyle(TableStyle([("ALIGN", (1,0), (-1,-1), "RIGHT"), ("FONTNAME", (1,0), (-1,1), "Helvetica"), ("FONTSIZE", (1,0), (-1,-1), 8.5), ("LINEABOVE", (1,2), (-1,2), 1, colors.HexColor(f"#{BURGUNDY}")), ("BACKGROUND", (1,2), (-1,2), colors.HexColor(f"#{BURGUNDY_LIGHT}")), ("TEXTCOLOR", (1,2), (-1,2), colors.HexColor(f"#{BURGUNDY}")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("TOPPADDING", (0,0), (-1,-1), 6), ("BOTTOMPADDING", (0,0), (-1,-1), 6)]))
    story += [totals, Spacer(1, 0.21*inch)]
    bank = Table([[Paragraph("<font color='#6F6870' size='7'><b>WIRE INSTRUCTIONS</b></font><br/>Beneficiary: Apex Strategic Services LLC<br/>Bank: First Continental Bank<br/>Routing: 021000089<br/>Account: Personal checking ending 4401", body), Paragraph("<font color='#6F6870' size='7'><b>CONTACT</b></font><br/>Billing questions: accounts@apexstrategic.example<br/>Engagement lead: Cameron Reed<br/>Vendor ID: pending assignment<br/>Tax ID: **-***4821", body)]], colWidths=[3.42*inch, 3.43*inch])
    bank.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#FFF8E8")), ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#E6CF9D")), ("INNERGRID", (0,0), (-1,-1), 0.5, colors.HexColor("#E6CF9D")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 11), ("RIGHTPADDING", (0,0), (-1,-1), 11), ("TOPPADDING", (0,0), (-1,-1), 9), ("BOTTOMPADDING", (0,0), (-1,-1), 9)]))
    story += [bank, Spacer(1, 0.18*inch), Paragraph("Thank you for your business. This document was generated electronically and does not require a signature.", small), Paragraph("Apex Strategic Services LLC · Confidential client billing document · Page 1 of 1", ParagraphStyle("foot", parent=small, alignment=TA_RIGHT, spaceBefore=5))]
    doc.build(story)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    add_csv()
    add_json()
    add_xlsx()
    add_docx()
    add_pdf()
    print(f"Generated upgraded fixtures in {OUT}")
